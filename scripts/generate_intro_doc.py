# -*- coding: utf-8 -*-
"""生成 X-RAG 知识库+指标管理+智能问答 产品介绍文档"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from datetime import date

# 品牌色
PRIMARY = RGBColor(0x4F, 0x46, 0xE5)      # Indigo-600
PRIMARY_LIGHT = RGBColor(0xE0, 0xE7, 0xFF)  # Indigo-100
ACCENT = RGBColor(0x05, 0x96, 0x69)        # Emerald-600
TEXT_DARK = RGBColor(0x1E, 0x29, 0x3B)     # Slate-800
TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)    # Slate-500
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_cell_shading(cell, hex_color: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def set_run_font(run, size=11, bold=False, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(text)
        set_run_font(run, size=18, bold=True, color=PRIMARY)
        # 下划线装饰
        border_p = doc.add_paragraph()
        border_p.paragraph_format.space_before = Pt(0)
        border_p.paragraph_format.space_after = Pt(14)
        run2 = border_p.add_run("━" * 42)
        set_run_font(run2, size=8, color=PRIMARY)
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(f"▎ {text}")
        set_run_font(run, size=14, bold=True, color=TEXT_DARK)
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"● {text}")
        set_run_font(run, size=12, bold=True, color=ACCENT)
    return p


def add_body(doc, text, indent=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    if indent:
        p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    set_run_font(run, size=11, color=TEXT_DARK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(3)
    if level > 0:
        p.paragraph_format.left_indent = Cm(0.8 * level)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=TEXT_DARK)
    return p


def add_feature_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # 表头
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=WHITE)
        set_cell_shading(cell, "4F46E5")

    # 数据行
    for ri, row in enumerate(rows):
        tr = table.rows[ri + 1]
        bg = "F8FAFC" if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=9.5, color=TEXT_DARK)
            set_cell_shading(cell, bg)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    doc.add_paragraph()  # 间距
    return table


def add_highlight_box(doc, title, content_lines):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "EEF2FF")
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(title)
    set_run_font(run, size=11, bold=True, color=PRIMARY)
    for line in content_lines:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        run2 = p2.add_run(f"  • {line}")
        set_run_font(run2, size=10, color=TEXT_DARK)
    doc.add_paragraph()


def create_cover_page(doc):
    # 顶部色块
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "4F46E5")
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(40)
    run = p.add_run("X-RAG 知识工作空间")
    set_run_font(run, size=28, bold=True, color=WHITE)

    doc.add_paragraph()
    doc.add_paragraph()

    # 副标题
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("知识库 · 指标管理 · 智能问数")
    set_run_font(run2, size=20, bold=True, color=PRIMARY)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run("企业级可信知识全生命周期管理平台")
    set_run_font(run3, size=13, color=TEXT_MUTED)

    doc.add_paragraph()
    doc.add_paragraph()

    # 核心卖点
    box = doc.add_table(rows=1, cols=1)
    bc = box.rows[0].cells[0]
    set_cell_shading(bc, "ECFDF5")
    bc.text = ""
    lines = [
        "以「指标治理 + 智能问数」为核心，打通指标口径、知识资产与业务问答",
        "从资料上传、AI 抽取、审核入库，到精准检索、可信问答、质量治理",
        "支持多角色权限隔离、来源溯源与开放 API 集成"
    ]
    for i, line in enumerate(lines):
        bp = bc.paragraphs[0] if i == 0 else bc.add_paragraph()
        bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        br = bp.add_run(line)
        set_run_font(br, size=11, color=ACCENT if i == 0 else TEXT_DARK)

    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    today_str = f"{date.today().year}-{date.today().month:02d}-{date.today().day:02d}"
    run4 = p4.add_run(f"文档版本：V1.0  |  生成日期：{today_str}")
    set_run_font(run4, size=10, color=TEXT_MUTED)

    doc.add_page_break()


def build_document():
    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # 默认字体
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(11)

    create_cover_page(doc)

    # ========== 目录 ==========
    add_heading_styled(doc, "目  录", level=1)
    toc_items = [
        "一、产品概述",
        "二、核心价值：指标管理 + 智能问数",
        "三、指标管理能力详解",
        "四、智能问数能力详解",
        "五、知识库管理能力",
        "六、知识生产与治理",
        "七、权限安全与开放集成",
        "八、系统架构与技术栈",
        "九、应用场景与价值",
        "附录：完整功能清单"
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=11, color=TEXT_DARK)
    doc.add_page_break()

    # ========== 一、产品概述 ==========
    add_heading_styled(doc, "一、产品概述", level=1)
    add_body(doc,
        "X-RAG 知识工作空间（x-rag-knowledge-platform）是一款面向企业的可信知识全生命周期管理平台。"
        "平台以「知识库」为底座、「指标管理」为治理核心、「智能问数」为业务出口，"
        "帮助企业将分散在制度文件、操作手册、指标口径文档、FAQ 等资料中的知识，"
        "转化为可搜索、可问答、可溯源、可治理的结构化知识资产。"
    )
    add_highlight_box(doc, "产品定位", [
        "可信知识中台：统一沉淀、审核、发布企业知识",
        "指标治理中心：管理指标口径、公式、数据来源与变更历史",
        "智能问数入口：自然语言提问，获取结构化、可追溯的答案",
        "质量治理闭环：持续体检、发现问题、推动修复与复审"
    ])

    add_heading_styled(doc, "平台八大功能模块", level=2)
    add_feature_table(doc,
        ["模块", "定位", "核心页面"],
        [
            ["工作台", "个人工作入口与运营总览", "总览 / 待办 / 最近使用 / 运营动态"],
            ["知识资产", "知识库、目录、标签、地图管理", "知识库 / 目录 / 标签 / 地图 / 版本"],
            ["知识生产", "资料入库与 AI 抽取流水线", "上传 / 任务 / 技能 / 工具库"],
            ["知识应用", "搜索、问答、场景与反馈", "智能问答 / 搜索 / 场景 / 收藏 / 反馈"],
            ["质量治理", "知识健康体检与问题闭环", "体检 / 问题 / 任务 / 冲突 / 复审 / 规则"],
            ["权限安全", "角色、密级、鉴权与审计", "概览 / 密级 / 角色 / SSO / 鉴权策略 / 审计"],
            ["开放集成", "API 与第三方系统对接", "API / 应用 / 输出适配 / Webhook / 日志"],
            ["系统设置", "平台参数与模型配置", "基础 / 通知 / 字典 / 模型 / 日志"],
        ],
        col_widths=[2.5, 5.5, 7.5]
    )

    # ========== 二、核心价值 ==========
    add_heading_styled(doc, "二、核心价值：指标管理 + 智能问数", level=1)
    add_body(doc,
        "在企业数据驱动决策的背景下，「指标口径不一致」「问指标答不准」「数据来源说不清」"
        "是普遍痛点。X-RAG 将指标知识纳入统一治理体系，并通过智能问数能力，"
        "让用户用自然语言即可获取权威、结构化、可溯源的指标解释与业务答案。"
    )

    add_heading_styled(doc, "双核驱动模型", level=2)
    add_feature_table(doc,
        ["能力维度", "解决什么问题", "典型用户问题"],
        [
            ["指标管理", "指标口径分散、版本混乱、变更无记录", "达标率怎么算？口径变更后历史报告怎么办？"],
            ["智能问数", "找知识难、答非所问、答案不可信", "数据共享审批流程是什么？某制度能否用于对外报告？"],
            ["知识联动", "指标与制度、流程知识割裂", "指标定义依据哪份制度？整改流程与指标验收如何衔接？"],
            ["治理闭环", "低质量知识、过期口径持续被引用", "哪些指标知识需要复审？冲突口径如何合并？"],
        ],
        col_widths=[3, 5.5, 7]
    )

    add_heading_styled(doc, "端到端价值链", level=2)
    steps = [
        "资料采集 → 上传 PDF/Word/Excel 等指标口径文档、制度文件",
        "智能抽取 → 财务指标抽取、表格抽取、QA 对生成等 23 种工具自动解析",
        "审核入库 → 人工审核 AI 候选知识，确认口径、公式、适用范围",
        "标签治理 → 领域/场景/知识类型（含「指标类」）五维标签精准归类",
        "精准检索 → 按指标类、可信等级、质量评分等多维过滤命中知识",
        "智能问数 → 选择「指标解释」等场景，自然语言提问获结构化答案",
        "反馈优化 → 用户反馈驱动知识缺口识别与治理任务派发"
    ]
    for i, step in enumerate(steps, 1):
        add_bullet(doc, f"步骤 {i}：{step}")

    doc.add_page_break()

    # ========== 三、指标管理 ==========
    add_heading_styled(doc, "三、指标管理能力详解", level=1)
    add_body(doc,
        "指标管理并非孤立的数据字典，而是将指标口径、计算公式、数据来源、适用场景、"
        "变更规则等知识化、版本化、标签化，纳入与企业制度、流程知识同一套治理体系。"
    )

    add_heading_styled(doc, "3.1 指标知识体系", level=2)
    add_feature_table(doc,
        ["功能项", "功能说明", "业务价值"],
        [
            ["指标类知识分类", "将指标口径文档、管理办法等标记为「指标类」知识", "与制度类、流程类、FAQ 类区分管理"],
            ["指标解释场景标签", "专属「指标解释」场景标签，参与检索与问答路由", "问指标类问题自动命中正确知识范围"],
            ["高价值指标知识", "如《指标口径管理办法》纳入高价值知识榜单", "优先推荐、优先命中、优先维护"],
            ["指标知识身份", "编码、来源文件、责任部门、版本、有效期等元数据", "每条指标知识可追溯、可问责"],
            ["指标口径一致性", "质量评分含「准确一致性」维度，检测口径冲突", "避免同一指标多套口径并存"],
        ],
        col_widths=[3.5, 6, 5.5]
    )

    add_heading_styled(doc, "3.2 指标知识生产", level=2)
    add_feature_table(doc,
        ["功能项", "功能说明", "技术能力"],
        [
            ["财务指标抽取", "从财报、经营分析文档中抽取金额、增长率、周期指标", "financial_indicator_extractor 工具"],
            ["表格抽取", "识别表头、单元格，解析指标定义表", "table_extractor 工具"],
            ["Excel/CSV 解析", "结构化解析指标清单、维度表", "excel_parser_tool"],
            ["QA 对生成", "从指标文档片段自动生成可审核问答对", "qa_pair_generator 工具"],
            ["摘要生成", "为指标知识生成摘要与适用场景说明", "summary_generator 工具"],
            ["抽取技能编排", "可视化编排 PDF/Word 解析 → 抽取 → 入库流水线", "ReactFlow 技能编排器"],
        ],
        col_widths=[3.5, 6, 5.5]
    )

    add_heading_styled(doc, "3.3 指标知识治理", level=2)
    add_feature_table(doc,
        ["功能项", "功能说明", "示例"],
        [
            ["口径统一任务", "发现口径不一致时创建治理任务", "「统一数据质量达标率口径」"],
            ["指标知识缺口", "统计高频未覆盖的指标问题", "「指标口径变更后历史报告如何处理？」"],
            ["冲突合并", "处理重复或冲突的指标定义", "多版本口径合并为权威版本"],
            ["复审下架", "过期或长期未更新的指标知识下架", "口径废止后自动提醒复审"],
            ["治理规则", "可配置规则自动发现指标知识问题", "摘要为空、来源不清晰、质量分偏低"],
        ],
        col_widths=[3.5, 6, 5.5]
    )

    add_heading_styled(doc, "3.4 指标检索与发现", level=2)
    add_bullet(doc, "知识搜索支持按「指标类」知识类型过滤")
    add_bullet(doc, "支持按领域标签（数据治理、数据质量、指标解释等）精准定位")
    add_bullet(doc, "搜索结果展示命中原因、可信等级（S/A 级）、质量评分、权限状态")
    add_bullet(doc, "知识地图展示「问题—标签—场景—知识」链路，含指标解释场景命中率")
    add_bullet(doc, "高频问题榜单含「数据质量达标率如何计算？」等典型指标问法")

    doc.add_page_break()

    # ========== 四、智能问数 ==========
    add_heading_styled(doc, "四、智能问数能力详解", level=1)
    add_body(doc,
        "智能问数是 X-RAG 面向业务用户的核心出口。用户以自然语言提问，"
        "系统自动完成意图识别、标签匹配、知识检索、权限校验与结构化答案生成，"
        "每一条答案均可溯源至权威知识来源。"
    )

    add_heading_styled(doc, "4.1 九大问答场景", level=2)
    add_feature_table(doc,
        ["场景", "适用问题类型", "答案模板"],
        [
            ["制度查询", "制度条款、版本、适用范围", "条款引用 + 制度依据"],
            ["流程咨询", "办理步骤、责任部门、审批节点", "流程说明 + 制度依据"],
            ["指标解释 ★", "指标口径、公式、数据来源、计算逻辑", "口径定义 + 公式说明 + 来源依据"],
            ["报告生成", "内部报告素材组织与引用", "结构化报告段落 + 引用清单"],
            ["方案编写", "方案素材与行业表达", "方案框架 + 参考知识"],
            ["标书辅助", "可对外引用知识与案例", "白名单知识 + 合规提示"],
            ["客服问答", "高频标准化问题", "标准答案 + FAQ 来源"],
            ["风险合规判断", "敏感、受限、输出风险识别", "合规结论 + 风险提示"],
            ["项目经验参考", "项目案例与专家经验复用", "案例摘要 + 经验要点"],
        ],
        col_widths=[3, 5.5, 6.5]
    )

    add_heading_styled(doc, "4.2 问数处理流程", level=2)
    add_feature_table(doc,
        ["阶段", "系统行为", "用户可见信息"],
        [
            ["问题理解", "识别意图、领域、标签、场景", "问题理解面板：意图/标签/场景/策略"],
            ["范围限定", "按场景限定检索知识库与标签范围", "检索范围：知识库列表 + 可信等级过滤"],
            ["知识命中", "向量检索 + 重排，匹配最相关知识", "命中知识列表：质量分/可信等级/命中原因"],
            ["权限校验", "校验当前用户查看/引用/输出权限", "权限状态：通过/受限/拒绝"],
            ["答案生成", "按场景模板生成结构化答案", "结论/步骤/责任部门/风险提示/来源依据"],
            ["来源溯源", "关联原始文件、章节、版本", "来源文件/章节/质量评分/权限状态"],
            ["反馈闭环", "收集有帮助/不准确/需补充等反馈", "反馈进入质量治理与知识优化"],
        ],
        col_widths=[2.5, 5.5, 7]
    )

    add_heading_styled(doc, "4.3 结构化答案要素", level=2)
    for item in ["结论：直接回答用户问题的核心要点",
                 "适用场景：明确该答案适用的业务情境",
                 "办理步骤：流程类问题的分步操作指引",
                 "责任部门：明确各环节责任主体",
                 "风险提示：合规、安全、输出风险预警",
                 "来源依据：引用制度文件、章节与版本信息"]:
        add_bullet(doc, item)

    add_heading_styled(doc, "4.4 指标问数专题能力", level=2)
    add_highlight_box(doc, "指标问数典型场景", [
        "「数据质量达标率如何计算？」→ 命中指标类知识，返回公式与口径说明",
        "「指标口径变更后历史报告如何处理？」→ 识别知识缺口，推动治理任务",
        "「某个指标能否用于对外报告？」→ 结合合规场景判断输出权限",
        "「达标率与合格率有什么区别？」→ 检索多指标知识，检测口径一致性"
    ])

    add_heading_styled(doc, "4.5 精准知识搜索", level=2)
    add_feature_table(doc,
        ["过滤维度", "可选值/说明"],
        [
            ["知识类型", "制度类 / 流程类 / FAQ类 / 指标类 / 案例类"],
            ["可信等级", "S级权威 / A级审核 / B级候选"],
            ["质量评分", "按质量分区间过滤"],
            ["知识状态", "已发布 / 已审核 / 草稿 / 已下架"],
            ["权限范围", "当前用户可查看 / 可引用 / 可输出"],
            ["场景标签", "指标解释 / 流程咨询 / 报告生成等"],
            ["输出标签", "可用于问答 / 报告 / 标书 / 对外输出"],
        ],
        col_widths=[4, 11]
    )

    doc.add_page_break()

    # ========== 五、知识库管理 ==========
    add_heading_styled(doc, "五、知识库管理能力", level=1)

    add_heading_styled(doc, "5.1 知识库管理", level=2)
    add_feature_table(doc,
        ["功能项", "详细说明"],
        [
            ["四级知识库体系", "个人 / 部门 / 企业 / 平台四级知识库分类管理"],
            ["知识库创建", "配置名称、颜色、类型、密级、描述"],
            ["运营指标看板", "条目数、待审核数、质量分、密级、状态一览"],
            ["角色隔离", "按用户角色过滤可见知识库范围"],
            ["预置知识库", "营销、售前、客户问答、员工制度、人事流程等 9 个演示库"],
        ],
        col_widths=[4, 11]
    )

    add_heading_styled(doc, "5.2 知识目录", level=2)
    add_bullet(doc, "知识库 → 分类层级目录树浏览")
    add_bullet(doc, "知识列表：标题、类型、标签、密级、质量分")
    add_bullet(doc, "知识详情抽屉：统一展示知识全貌")
    add_bullet(doc, "批量操作：打标、调密级、移动、新建分类")

    add_heading_styled(doc, "5.3 五维标签体系", level=2)
    add_feature_table(doc,
        ["标签类型", "用途", "参与环节"],
        [
            ["领域标签", "限定知识所属行业领域", "检索过滤 / 问答范围 / 知识地图"],
            ["场景标签", "标识知识适用业务场景（含指标解释）", "答案模板选择 / 场景路由"],
            ["知识类型标签", "区分制度/流程/FAQ/指标/案例", "检索分类 / 问答策略"],
            ["权限标签", "内部/敏感/受限等访问控制", "权限校验 / 输出控制"],
            ["输出标签", "可用于问答/报告/标书/对外", "输出权限判断"],
        ],
        col_widths=[3, 5, 7]
    )

    add_heading_styled(doc, "5.4 知识地图", level=2)
    add_bullet(doc, "问题—标签—场景—知识链路可视化，含命中率与满意度")
    add_bullet(doc, "知识关系图：包含/引用/关联/同源关系（ReactFlow）")
    add_bullet(doc, "类型×知识库矩阵：发现知识盲区与质量薄弱区")
    add_bullet(doc, "审批时效统计：通过率、平均时长、一次通过率")

    add_heading_styled(doc, "5.5 版本记录", level=2)
    add_bullet(doc, "知识条目版本历史追踪")
    add_bullet(doc, "变更说明、编辑人、编辑时间记录")
    add_bullet(doc, "当前版本 / 历史版本状态管理")

    # ========== 六、知识生产与治理 ==========
    add_heading_styled(doc, "六、知识生产与治理", level=1)

    add_heading_styled(doc, "6.1 知识生产", level=2)
    add_feature_table(doc,
        ["功能项", "详细说明"],
        [
            ["文件上传向导", "4 步向导：选文件 → 选技能 → 配置类型/知识库/密级 → 提交"],
            ["支持格式", "PDF / Word / Excel / PPT / TXT / MD / JSON / 网页 / 音频"],
            ["抽取任务管理", "任务列表、进度追踪、状态统计"],
            ["抽取技能编排", "ReactFlow 可视化编排解析流水线"],
            ["工具库", "23 个内置工具：解析/清洗/切分/抽取/知识构建"],
            ["审核工作台", "AI 候选知识审核：通过/驳回/编辑/批量操作"],
        ],
        col_widths=[4, 11]
    )

    add_heading_styled(doc, "6.2 质量治理", level=2)
    add_feature_table(doc,
        ["功能项", "详细说明"],
        [
            ["知识体检", "八维度质量评分：权威性/完整性/时效性/一致性等"],
            ["问题清单", "摘要缺失/缺少变体/标签缺失/来源不清/质量偏低"],
            ["治理任务", "待处理/处理中/已完成/逾期任务管理"],
            ["冲突合并", "重复/冲突知识识别与合并处理"],
            ["复审下架", "到期/过期/长期未更新/已下架生命周期管理"],
            ["治理规则", "可配置自动检测规则"],
            ["价值分类", "按使用频次、反馈、场景覆盖评估知识价值"],
        ],
        col_widths=[4, 11]
    )

    doc.add_page_break()

    # ========== 七、权限与集成 ==========
    add_heading_styled(doc, "七、权限安全与开放集成", level=1)

    add_heading_styled(doc, "7.1 权限安全", level=2)
    add_feature_table(doc,
        ["功能项", "详细说明"],
        [
            ["角色权限 RBAC", "超级管理员/企业管理员/业务管理员等多角色"],
            ["模块级权限", "按模块控制 workspace/assets/production 等访问"],
            ["知识库级隔离", "accessibleKnowledgeBases 范围控制"],
            ["密级配置", "L1-L4 四级密级体系"],
            ["访问标签", "细粒度访问控制标签"],
            ["知识鉴权策略", "查看/引用/下载/编辑/审核/智能体调用/对外输出"],
            ["SSO 配置", "单点登录集成"],
            ["安全告警", "异常访问与违规调用告警"],
            ["操作审计", "全量操作日志追溯"],
        ],
        col_widths=[4, 11]
    )

    add_heading_styled(doc, "7.2 开放集成", level=2)
    add_feature_table(doc,
        ["功能项", "详细说明"],
        [
            ["API 管理", "应用列表、API Key、限流策略"],
            ["应用接入", "第三方系统接入配置"],
            ["输出适配", "Output Profile、转换规则、格式预览、导出历史"],
            ["Webhook", "知识审核/更新/下架事件通知"],
            ["调用日志", "Token 用量、耗时、错误信息统计"],
        ],
        col_widths=[4, 11]
    )

    add_heading_styled(doc, "7.3 系统设置", level=2)
    add_bullet(doc, "基础设置：平台名、企业名、体验偏好")
    add_bullet(doc, "通知设置：待审核/SLA/治理/到期/安全/API 异常通知")
    add_bullet(doc, "字典配置：知识类型、密级、反馈类型、治理问题、应用类型")
    add_bullet(doc, "模型配置：问答 / Embedding / 重排 / 摘要模型")
    add_bullet(doc, "系统日志：平台运行日志查看")

    # ========== 八、架构 ==========
    add_heading_styled(doc, "八、系统架构与技术栈", level=1)

    add_heading_styled(doc, "8.1 技术栈", level=2)
    add_feature_table(doc,
        ["层次", "技术选型"],
        [
            ["前端框架", "React 18 + TypeScript 5.9 + Vite 5.4"],
            ["路由与状态", "React Router 6 + Zustand 4"],
            ["UI 组件", "Radix UI + Tailwind CSS 3.4 + Lucide Icons"],
            ["可视化", "Recharts（图表）+ ReactFlow（关系图/技能编排）"],
            ["认证授权", "角色演示登录 + RBAC 模块/路径/知识库三级权限"],
            ["部署方式", "静态 SPA，支持 Nginx / IIS / 对象存储部署"],
        ],
        col_widths=[4, 11]
    )

    add_heading_styled(doc, "8.2 系统架构", level=2)
    arch_lines = [
        "┌─────────────────────────────────────────────────┐",
        "│           浏览器 SPA（React + Vite）              │",
        "├─────────────────────────────────────────────────┤",
        "│  展示层：8 大模块 × 40+ 页面                       │",
        "│  状态层：Zustand 全局状态 + 角色认证               │",
        "│  业务层：权限控制 / 角色过滤 / 导航配置             │",
        "│  数据层：知识库 / 指标 / 问答 Mock 数据             │",
        "├─────────────────────────────────────────────────┤",
        "│  接入层（可扩展）                                  │",
        "│    → RAG 检索服务（Embedding + 重排）             │",
        "│    → LLM 问答服务（结构化答案生成）                │",
        "│    → 指标元数据服务（口径/维度/血缘）              │",
        "│    → 权限与审计服务                               │",
        "└─────────────────────────────────────────────────┘",
    ]
    for line in arch_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(line)
        set_run_font(run, size=9, color=TEXT_MUTED, name="Consolas")

    add_body(doc,
        "当前版本为前端演示系统，完整展示了产品功能蓝图与交互体验。"
        "接入真实后端后，可无缝对接 RAG 检索引擎、大语言模型、指标元数据库与企业 SSO 等基础设施。"
    )

    # ========== 九、应用场景 ==========
    add_heading_styled(doc, "九、应用场景与价值", level=1)
    add_feature_table(doc,
        ["应用场景", "指标管理价值", "智能问数价值"],
        [
            ["经营分析", "统一营收、利润率等核心指标口径", "「毛利率如何计算？」即时获权威答案"],
            ["数据治理", "沉淀数据质量、资产等治理指标知识", "「达标率整改流程？」关联流程知识"],
            ["合规审查", "管理对外输出指标的合规边界", "「能否用于对外报告？」自动合规判断"],
            ["部门协同", "跨部门指标口径一致性治理", "「A部门与B部门口径差异？」冲突检测"],
            ["新人培训", "指标知识库作为培训教材", "自然语言提问快速上手业务指标"],
            ["智能报告", "报告引用指标知识自动标注来源", "「生成月度质量报告」结构化输出"],
        ],
        col_widths=[3, 5.5, 6.5]
    )

    add_highlight_box(doc, "量化价值指标（平台运营看板）", [
        "知识可信率：87.6%（已通过来源、版本和质量校验）",
        "问答命中准确率：91.3%（用户问题命中正确知识和场景）",
        "高价值知识数：326 条（高频使用且反馈稳定）",
        "知识缺口数：27 个（驱动持续知识补全）"
    ])

    doc.add_page_break()

    # ========== 附录 ==========
    add_heading_styled(doc, "附录：完整功能清单", level=1)
    add_body(doc, "以下为 X-RAG 平台全部功能点的完整清单，按模块—页面—功能三级组织。")

    appendix_data = [
        ("工作台", [
            ("总览", ["精准治理指标看板（8项）", "高价值知识 Top10", "高频问题 Top10", "知识缺口 Top10", "快捷入口（搜知识/问AI/上传/审核/体检/建库）"]),
            ("我的待办", ["待审核任务", "待治理任务", "到期复审提醒"]),
            ("最近使用", ["最近访问知识记录"]),
            ("运营动态", ["平台活动流/变更通知"]),
        ]),
        ("知识资产", [
            ("知识库", ["四级分类展示", "新建知识库", "知识库卡片指标", "角色可见范围过滤"]),
            ("知识目录", ["目录树浏览", "知识列表", "详情抽屉", "批量打标/调密级/移动"]),
            ("标签体系", ["五类标签管理", "标签统计", "新增/合并/纠偏", "关联知识查看"]),
            ("知识地图", ["问题链路图", "知识关系图", "类型矩阵", "审批时效", "知识条目浏览"]),
            ("版本记录", ["版本历史", "变更说明", "版本状态"]),
        ]),
        ("知识生产", [
            ("文件上传", ["4步上传向导", "多格式支持", "技能选择", "密级配置"]),
            ("抽取任务", ["任务列表", "进度追踪", "状态统计"]),
            ("抽取技能", ["可视化编排", "流水线预览"]),
            ("工具库", ["23个内置工具", "工具详情", "运行指标"]),
            ("审核工作台", ["候选知识审核", "通过/驳回/编辑", "批量操作"]),
        ]),
        ("知识应用 ★", [
            ("智能问答", ["9大场景选择", "推荐问题", "问题理解面板", "结构化答案", "命中知识", "来源溯源", "用户反馈"]),
            ("知识搜索", ["精准检索", "多维过滤（含指标类）", "命中原因", "详情抽屉"]),
            ("应用场景", ["场景配置", "调用量/命中率", "知识范围"]),
            ("收藏与常用", ["收藏知识", "常用搜索词", "常用知识库"]),
            ("用户反馈", ["反馈类型", "转治理任务"]),
        ]),
        ("质量治理", [
            ("知识体检", ["八维评分", "低质量清单", "价值分类", "整改任务"]),
            ("问题清单", ["5类问题类型", "问题详情"]),
            ("治理任务", ["任务状态管理", "优先级/截止日期"]),
            ("冲突合并", ["重复识别", "合并处理"]),
            ("复审下架", ["到期管理", "下架流程"]),
            ("治理规则", ["规则配置", "自动检测"]),
        ]),
        ("权限安全", [
            ("权限概览", ["权限统计看板"]),
            ("密级配置", ["L1-L4密级"]),
            ("访问标签", ["标签管理"]),
            ("用户与角色", ["角色分配", "权限配置"]),
            ("SSO配置", ["单点登录"]),
            ("安全告警", ["异常告警"]),
            ("知识鉴权策略", ["7种权限维度"]),
            ("操作审计", ["审计日志"]),
        ]),
        ("开放集成", [
            ("API管理", ["应用/Key/限流"]),
            ("应用接入", ["接入配置"]),
            ("输出适配", ["Profile/规则/预览/历史"]),
            ("Webhook", ["事件通知"]),
            ("调用日志", ["用量/耗时/错误"]),
        ]),
        ("系统设置", [
            ("基础设置", ["平台参数"]),
            ("通知设置", ["6类通知"]),
            ("字典配置", ["5类字典"]),
            ("模型配置", ["4类模型"]),
            ("系统日志", ["运行日志"]),
        ]),
    ]

    for module_name, pages in appendix_data:
        add_heading_styled(doc, module_name, level=2)
        for page_name, features in pages:
            add_heading_styled(doc, page_name, level=3)
            for feat in features:
                add_bullet(doc, feat)

    # 页脚
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("— 文档结束 —")
    set_run_font(run, size=10, color=TEXT_MUTED)

    return doc


if __name__ == "__main__":
    output_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "docs")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "X-RAG-知识库-指标管理-智能问数-产品介绍.docx")

    doc = build_document()
    doc.save(output_path)
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    print(output_path)
